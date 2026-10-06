from calendar import monthrange
from collections import defaultdict
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from .. import models
from ..database import get_db
from ..deps import exiger_admin

router = APIRouter(prefix="/exports", tags=["Exports"])

FORMAT_MONTANT = "#,##0.000"
FORMAT_DATE = "DD/MM/YYYY"
ENTETE_FILL = PatternFill("solid", fgColor="1F3864")
ENTETE_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)

LIBELLES_STATUT = {"payee": "Payée", "impayee": "Impayée", "partielle": "Partielle"}
LIBELLES_CATEGORIE = {
    "salaire_chauffeur": "Salaire chauffeur",
    "cnss": "CNSS",
    "carburant": "Carburant",
    "entretien": "Entretien",
    "assurance": "Assurance",
    "taxe": "Taxe",
    "autre": "Autre",
}
TRANCHES = ["0-30 j", "31-60 j", "61-90 j", "> 90 j"]


def _bornes(date_du: date | None, date_au: date | None) -> tuple[date, date]:
    """Par défaut : le mois civil en cours."""
    if date_du and date_au:
        return date_du, date_au
    today = date.today()
    return (
        date(today.year, today.month, 1),
        date(today.year, today.month, monthrange(today.year, today.month)[1]),
    )


def _valeur(x) -> str:
    return x.value if hasattr(x, "value") else str(x)


def _entetes(ws, colonnes: list[str], ligne: int = 1):
    for i, titre in enumerate(colonnes, start=1):
        c = ws.cell(row=ligne, column=i, value=titre)
        c.fill = ENTETE_FILL
        c.font = ENTETE_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Chaîne (et non ws.cell) : accéder à une cellule la crée et décale ws.append.
    ws.freeze_panes = f"A{ligne + 1}"


def _formater(ws, cols_montant: list[int] = (), cols_date: list[int] = (), a_partir_de: int = 2):
    for row in ws.iter_rows(min_row=a_partir_de):
        for c in row:
            if c.column in cols_montant:
                c.number_format = FORMAT_MONTANT
            elif c.column in cols_date:
                c.number_format = FORMAT_DATE


def _largeurs(ws, largeurs: list[int]):
    for i, w in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _ligne_total(ws, libelle: str, cols_somme: list[int], premiere: int, derniere: int):
    """Ligne de total avec formules SUM, pour que le comptable puisse vérifier."""
    r = derniere + 1
    ws.cell(row=r, column=1, value=libelle).font = TOTAL_FONT
    for col in cols_somme:
        lettre = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({lettre}{premiere}:{lettre}{derniere})")
        c.font = TOTAL_FONT
        c.number_format = FORMAT_MONTANT


@router.get("/comptable", dependencies=[Depends(exiger_admin)])
def export_comptable(
    date_du: date | None = None,
    date_au: date | None = None,
    db: Session = Depends(get_db),
):
    """
    Export Excel pour le comptable, sur une période (mois en cours par défaut).

    La date de facture retenue est `date_fin`, la même que celle imprimée sur
    le PDF. Les factures de location (client en texte libre) sont fusionnées
    avec les factures classiques ; elles n'ont ni timbre ni matricule fiscal.
    """
    du, au = _bornes(date_du, date_au)
    if du > au:
        raise HTTPException(400, "La date de début doit précéder la date de fin.")

    # ---- Données -------------------------------------------------------
    factures = (
        db.query(models.Facture)
        .options(joinedload(models.Facture.client))
        .filter(models.Facture.date_fin >= du, models.Facture.date_fin <= au)
        .order_by(models.Facture.date_fin, models.Facture.numero_facture)
        .all()
    )
    factures_loc = (
        db.query(models.FactureLocation)
        .filter(models.FactureLocation.date_fin >= du, models.FactureLocation.date_fin <= au)
        .order_by(models.FactureLocation.date_fin, models.FactureLocation.numero_facture)
        .all()
    )

    # Ligne normalisée commune aux deux types de factures.
    lignes = []
    for f in factures:
        lignes.append({
            "date": f.date_fin, "numero": f.numero_facture, "type": "Navettes",
            "client": f.client.nom_societe, "mf": f.client.matricule_fiscal or "",
            "ht": float(f.montant_ht), "taux": float(f.taux_tva), "tva": float(f.montant_tva),
            "timbre": float(f.timbre or 0), "ttc": float(f.montant_ttc),
            "statut": _valeur(f.statut), "date_paiement": f.date_paiement,
        })
    for f in factures_loc:
        lignes.append({
            "date": f.date_fin, "numero": f.numero_facture, "type": "Location",
            "client": f.client, "mf": "",
            "ht": float(f.montant_ht), "taux": float(f.taux_tva), "tva": float(f.montant_tva),
            "timbre": 0.0, "ttc": float(f.montant_ttc),
            "statut": _valeur(f.statut), "date_paiement": f.date_paiement,
        })
    lignes.sort(key=lambda l: (l["date"], l["numero"]))

    wb = Workbook()

    # ---- 1. Journal des ventes ----------------------------------------
    ws = wb.active
    ws.title = "Journal des ventes"
    _entetes(ws, ["Date", "N° facture", "Type", "Client", "Matricule fiscal",
                  "Montant HT", "Taux TVA %", "TVA", "Timbre", "Total TTC", "Statut", "Date paiement"])
    for l in lignes:
        ws.append([l["date"], l["numero"], l["type"], l["client"], l["mf"],
                   l["ht"], l["taux"], l["tva"], l["timbre"], l["ttc"],
                   LIBELLES_STATUT.get(l["statut"], l["statut"]), l["date_paiement"]])
    n = len(lignes)
    _formater(ws, cols_montant=[6, 8, 9, 10], cols_date=[1, 12])
    if n:
        _ligne_total(ws, "TOTAL", [6, 8, 9, 10], 2, n + 1)
    _largeurs(ws, [12, 16, 11, 34, 20, 14, 11, 12, 10, 14, 11, 14])
    ws.auto_filter.ref = f"A1:L{max(n + 1, 2)}"

    # ---- 2. Récap TVA --------------------------------------------------
    ws = wb.create_sheet("Récap TVA")
    _entetes(ws, ["Taux TVA %", "Nb factures", "Base HT", "TVA collectée", "Timbres", "Total TTC"])
    par_taux = defaultdict(lambda: {"nb": 0, "ht": 0.0, "tva": 0.0, "timbre": 0.0, "ttc": 0.0})
    for l in lignes:
        g = par_taux[l["taux"]]
        g["nb"] += 1
        g["ht"] += l["ht"]
        g["tva"] += l["tva"]
        g["timbre"] += l["timbre"]
        g["ttc"] += l["ttc"]
    for taux in sorted(par_taux):
        g = par_taux[taux]
        ws.append([taux, g["nb"], round(g["ht"], 3), round(g["tva"], 3),
                   round(g["timbre"], 3), round(g["ttc"], 3)])
    k = len(par_taux)
    _formater(ws, cols_montant=[3, 4, 5, 6])
    if k:
        _ligne_total(ws, "TOTAL", [2, 3, 4, 5, 6], 2, k + 1)
        ws.cell(row=k + 2, column=2).number_format = "0"
    ws.cell(row=k + 4, column=1,
            value="Le timbre fiscal est hors base TVA : TTC = HT + TVA + timbre.").font = Font(italic=True)
    _largeurs(ws, [12, 12, 16, 16, 12, 16])

    # ---- 3. Encaissements ---------------------------------------------
    ws = wb.create_sheet("Encaissements")
    _entetes(ws, ["Date paiement", "N° facture", "Client", "Total TTC encaissé"])
    payees = sorted(
        (l for l in _toutes_factures(db) if l["statut"] == "payee"
         and l["date_paiement"] and du <= l["date_paiement"] <= au),
        key=lambda l: (l["date_paiement"], l["numero"]),
    )
    for l in payees:
        ws.append([l["date_paiement"], l["numero"], l["client"], l["ttc"]])
    p = len(payees)
    _formater(ws, cols_montant=[4], cols_date=[1])
    if p:
        _ligne_total(ws, "TOTAL ENCAISSÉ", [4], 2, p + 1)
    _largeurs(ws, [15, 16, 34, 18])

    # ---- 4. Balance âgée ----------------------------------------------
    # Situation à la date de fin de période. Il n'existe pas encore de table de
    # paiements : une facture « partielle » est comptée pour son TTC complet.
    ws = wb.create_sheet("Balance âgée")
    _entetes(ws, ["Client", *TRANCHES, "Total dû"])
    balance = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for l in _toutes_factures(db):
        if l["date"] > au or l["statut"] == "payee":
            continue
        age = (au - l["date"]).days
        idx = 0 if age <= 30 else 1 if age <= 60 else 2 if age <= 90 else 3
        balance[l["client"]][idx] += l["ttc"]
    clients = sorted(balance)
    for c in clients:
        t = balance[c]
        ws.append([c, *[round(v, 3) for v in t], None])
    b = len(clients)
    for r in range(2, b + 2):
        ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})")
    _formater(ws, cols_montant=[2, 3, 4, 5, 6])
    if b:
        _ligne_total(ws, "TOTAL", [2, 3, 4, 5, 6], 2, b + 1)
    ws.cell(row=b + 4, column=1,
            value=f"Impayés et partielles au {au.strftime('%d/%m/%Y')} (ancienneté depuis la date de facture). "
                  "Les factures partielles sont comptées pour leur TTC complet.").font = Font(italic=True)
    _largeurs(ws, [34, 14, 14, 14, 14, 16])

    # ---- 5. Dépenses ---------------------------------------------------
    ws = wb.create_sheet("Dépenses")
    _entetes(ws, ["Date", "Catégorie", "Description", "Véhicule", "Chauffeur", "Transporteur", "Montant"])
    depenses = (
        db.query(models.Depense)
        .options(joinedload(models.Depense.vehicule), joinedload(models.Depense.chauffeur),
                 joinedload(models.Depense.transporteur))
        .filter(models.Depense.date >= du, models.Depense.date <= au)
        .order_by(models.Depense.date, models.Depense.id)
        .all()
    )
    par_cat = defaultdict(float)
    for d in depenses:
        cat = _valeur(d.categorie)
        par_cat[cat] += float(d.montant)
        ws.append([
            d.date, LIBELLES_CATEGORIE.get(cat, cat), d.description or "",
            d.vehicule.matricule if d.vehicule else "",
            f"{d.chauffeur.prenom} {d.chauffeur.nom}" if d.chauffeur else "",
            d.transporteur.nom_agence if d.transporteur else "",
            float(d.montant),
        ])
    m = len(depenses)
    _formater(ws, cols_montant=[7], cols_date=[1])
    if m:
        _ligne_total(ws, "TOTAL", [7], 2, m + 1)
    # Totaux par catégorie, sous le détail.
    r = m + 4
    ws.cell(row=r, column=1, value="Totaux par catégorie").font = TOTAL_FONT
    for cat in sorted(par_cat):
        r += 1
        ws.cell(row=r, column=2, value=LIBELLES_CATEGORIE.get(cat, cat))
        c = ws.cell(row=r, column=7, value=round(par_cat[cat], 3))
        c.number_format = FORMAT_MONTANT
    _largeurs(ws, [12, 20, 40, 16, 24, 22, 14])

    buf = BytesIO()
    wb.save(buf)
    nom = f"export_comptable_{du.isoformat()}_{au.isoformat()}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nom}"'},
    )


def _toutes_factures(db: Session) -> list[dict]:
    """Toutes les factures (navettes + location), sous forme normalisée."""
    out = []
    for f in db.query(models.Facture).options(joinedload(models.Facture.client)).all():
        out.append({"date": f.date_fin, "numero": f.numero_facture, "client": f.client.nom_societe,
                    "ttc": float(f.montant_ttc), "statut": _valeur(f.statut),
                    "date_paiement": f.date_paiement})
    for f in db.query(models.FactureLocation).all():
        out.append({"date": f.date_fin, "numero": f.numero_facture, "client": f.client,
                    "ttc": float(f.montant_ttc), "statut": _valeur(f.statut),
                    "date_paiement": f.date_paiement})
    return out
