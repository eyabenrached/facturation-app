"""Exports du journal comptable : Excel (openpyxl) et PDF (reportlab)."""
import io
from datetime import date
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .compta_plan import TYPES_OPERATION

NAVY = colors.HexColor("#15305a")
GOLD = colors.HexColor("#b8902e")
ENTETES = ["Date", "N° pièce", "Libellé", "Compte", "Compte auxiliaire", "Débit", "Crédit", "Référence", "Type d'opération"]


def _fmt_date(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _fmt(v) -> str:
    return "" if not v else f"{float(v):,.3f}".replace(",", " ")


def _lignes_plates(ecritures):
    """Une ligne de tableau par ligne d'écriture (les infos d'en-tête sont
    répétées sur chaque ligne pour que l'export Excel reste filtrable)."""
    for e in ecritures:
        for l in e.lignes:
            yield [
                e.date, e.numero_piece, l.libelle or e.libelle, f"{l.compte.numero} {l.compte.libelle}",
                l.compte_auxiliaire or "", float(l.debit or 0), float(l.credit or 0), e.reference or "",
                TYPES_OPERATION.get(e.type_operation, e.type_operation),
            ]


def journal_xlsx(ecritures, titre_periode: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    ws["A1"] = "Journal comptable"
    ws["A1"].font = Font(bold=True, size=14, color="15305A")
    ws["A2"] = titre_periode
    ws.append([])
    ws.append(ENTETES)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="15305A")
        c.alignment = Alignment(horizontal="center", vertical="center")
    total_d = total_c = 0.0
    for ligne in _lignes_plates(ecritures):
        ws.append(ligne)
        r = ws.max_row
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        for col in (6, 7):
            ws.cell(r, col).number_format = "#,##0.000"
        total_d += ligne[5]
        total_c += ligne[6]
    ws.append(["", "", "TOTAL", "", "", round(total_d, 3), round(total_c, 3), "", ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.number_format = "#,##0.000"
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{ws.max_row - 1}"
    for i, largeur in enumerate([12, 20, 42, 34, 28, 14, 14, 20, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def journal_pdf(ecritures, titre_periode: str, societe: str = "EURAFR TOURS") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Journal comptable")
    base = getSampleStyleSheet()
    titre = ParagraphStyle("t", parent=base["Heading1"], textColor=NAVY, fontSize=15, spaceAfter=2)
    sous = ParagraphStyle("s", parent=base["Normal"], textColor=colors.grey, fontSize=9)
    cell = ParagraphStyle("c", parent=base["Normal"], fontSize=7.5, leading=9)

    data = [ENTETES]
    total_d = total_c = 0.0
    for l in _lignes_plates(ecritures):
        total_d += l[5]
        total_c += l[6]
        data.append([_fmt_date(l[0]), l[1], Paragraph(escape(l[2]), cell), Paragraph(escape(l[3]), cell), Paragraph(escape(l[4]), cell), _fmt(l[5]), _fmt(l[6]), l[7], Paragraph(escape(l[8]), cell)])
    data.append(["", "", "TOTAL", "", "", _fmt(total_d), _fmt(total_c), "", ""])

    t = Table(data, repeatRows=1, colWidths=[18 * mm, 26 * mm, 50 * mm, 48 * mm, 32 * mm, 22 * mm, 22 * mm, 24 * mm, 25 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (5, 0), (6, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f6f5f1")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e6e3da")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 1, GOLD),
    ]))

    def pied(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(10 * mm, 8 * mm, f"{societe} — Journal comptable — édité le {_fmt_date(date.today())}")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build([Paragraph(f"{societe} — Journal comptable", titre), Paragraph(titre_periode, sous), Spacer(1, 4 * mm), t],
              onFirstPage=pied, onLaterPages=pied)
    return buf.getvalue()


# ----------------------------------------------------------- relevé / relance clients
def _style_doc():
    base = getSampleStyleSheet()
    return {
        "titre": ParagraphStyle("t2", parent=base["Heading1"], textColor=NAVY, fontSize=16, spaceAfter=2),
        "sous": ParagraphStyle("s2", parent=base["Normal"], textColor=colors.grey, fontSize=9),
        "corps": ParagraphStyle("b2", parent=base["Normal"], fontSize=10, leading=14),
        "cell": ParagraphStyle("c2", parent=base["Normal"], fontSize=8.5, leading=10.5),
    }


def _tableau(data, col_widths, num_cols, total_ligne=True):
    t = Table(data, repeatRows=1, colWidths=col_widths)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2 if total_ligne else -1), [colors.white, colors.HexColor("#f6f5f1")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#e6e3da")),
    ]
    for c in num_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    if total_ligne:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 1, GOLD)]
    t.setStyle(TableStyle(style))
    return t


def releve_client_pdf(rel: dict, societe: str = "EURAFR TOURS") -> bytes:
    st = _style_doc()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Relevé de compte — {rel['client']}")
    periode = "Toutes périodes"
    if rel["du"] or rel["au"]:
        periode = f"Période du {_fmt_date(rel['du']) if rel['du'] else '…'} au {_fmt_date(rel['au']) if rel['au'] else '…'}"
    data = [["Date", "Pièce", "Libellé", "Débit", "Crédit", "Solde"]]
    if rel["du"]:
        data.append(["", "", "Solde à l'ouverture", "", "", _fmt(rel["solde_ouverture"]) or "0.000"])
    for l in rel["lignes"]:
        data.append([_fmt_date(l["date"]), Paragraph(escape(l["piece"]), st["cell"]), Paragraph(escape(l["libelle"]), st["cell"]),
                     _fmt(l["debit"]), _fmt(l["credit"]), _fmt(l["solde"]) or "0.000"])
    data.append(["", "", "SOLDE À PAYER", "", "", _fmt(rel["solde_final"]) or "0.000"])
    elems = [
        Paragraph(f"{escape(societe)} — Relevé de compte client", st["titre"]),
        Paragraph(f"Client : <b>{escape(rel['client'])}</b> — {periode} — édité le {_fmt_date(date.today())}", st["sous"]),
        Spacer(1, 5 * mm),
        _tableau(data, [22 * mm, 36 * mm, 56 * mm, 22 * mm, 22 * mm, 22 * mm], [3, 4, 5]),
        Spacer(1, 6 * mm),
        Paragraph("Montants en dinars tunisiens (TND), TVA et timbre fiscal inclus.", st["sous"]),
    ]
    doc.build(elems)
    return buf.getvalue()


def relance_client_pdf(client: str, factures: list[dict], delai_jours: int, societe: str = "EURAFR TOURS") -> bytes:
    st = _style_doc()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
                            title=f"Relance — {client}")
    total = sum(float(f["reste"]) for f in factures)
    data = [["Facture", "Date", "Échéance", "Total TTC", "Déjà réglé", "Reste dû", "Retard"]]
    for f in factures:
        data.append([Paragraph(escape(f["numero"]), st["cell"]), _fmt_date(f["date"]), _fmt_date(f["echeance"]), _fmt(f["ttc"]),
                     _fmt(f["paye"]) or "0.000", _fmt(f["reste"]), f"{f['jours_retard']} j" if f["jours_retard"] else "—"])
    data.append(["TOTAL DÛ", "", "", "", "", _fmt(total), ""])
    elems = [
        Paragraph(escape(societe), st["titre"]),
        Paragraph(f"Le {_fmt_date(date.today())}", st["sous"]),
        Spacer(1, 8 * mm),
        Paragraph(f"À l'attention de : <b>{escape(client)}</b>", st["corps"]),
        Spacer(1, 4 * mm),
        Paragraph("<b>Objet : relance — factures en attente de règlement</b>", st["corps"]),
        Spacer(1, 4 * mm),
        Paragraph("Madame, Monsieur,", st["corps"]),
        Spacer(1, 3 * mm),
        Paragraph(
            f"Sauf erreur ou omission de notre part, les factures ci-dessous, payables sous {delai_jours} jours, "
            "n'ont pas encore été réglées intégralement. Nous vous remercions de bien vouloir procéder à leur règlement "
            "dans les meilleurs délais, ou de nous contacter si un paiement a déjà été effectué.", st["corps"]),
        Spacer(1, 5 * mm),
        _tableau(data, [34 * mm, 20 * mm, 20 * mm, 24 * mm, 24 * mm, 24 * mm, 14 * mm], [3, 4, 5, 6]),
        Spacer(1, 6 * mm),
        Paragraph("Montants en dinars tunisiens (TND).", st["sous"]),
        Spacer(1, 6 * mm),
        Paragraph("Nous vous prions d'agréer, Madame, Monsieur, l'expression de nos salutations distinguées.", st["corps"]),
        Spacer(1, 12 * mm),
        Paragraph(escape(societe), st["corps"]),
    ]
    doc.build(elems)
    return buf.getvalue()
